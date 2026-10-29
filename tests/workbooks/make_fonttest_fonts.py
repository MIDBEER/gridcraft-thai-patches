import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

out = sys.argv[1]
wb = Workbook()
ws = wb.active
ws.title = "fonttest2"

fonts = ["Calibri", "Tahoma", "Leelawadee UI", "TH Sarabun New", "Angsana New", "Cordia New", "Times New Roman", "Courier New"]
widths = {"A": 20, "B": 22, "C": 22, "D": 40}
for k, v in widths.items():
    ws.column_dimensions[k].width = v

hdr = ["font", "Latin only", "Thai only", "Thai + digits"]
for j, t in enumerate(hdr):
    c = ws.cell(row=1, column=j + 1, value=t)
    c.font = Font(name="Tahoma", size=11, bold=True)

for i, f in enumerate(fonts, start=2):
    ws.row_dimensions[i].height = 34
    ws.cell(row=i, column=1, value=f).font = Font(name="Tahoma", size=11)
    vals = ["Room 101 Hello 2569", "สวัสดีครับ ห้องเรียน", "ห้อง 101 ชั้น 2 จำนวน 35 คน 2569"]
    for j, v in enumerate(vals, start=2):
        c = ws.cell(row=i, column=j, value=v)
        c.font = Font(name=f, size=14)
        c.alignment = Alignment(vertical="center")

wb.save(out)
print("saved", out)
