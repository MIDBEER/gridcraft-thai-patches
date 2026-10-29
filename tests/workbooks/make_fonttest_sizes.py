import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

out = sys.argv[1]
wb = Workbook()
ws = wb.active
ws.title = "fonttest3"

for k, v in {"A": 22, "B": 36, "C": 36}.items():
    ws.column_dimensions[k].width = v

for j, t in enumerate(["font / size", "Latin only", "Thai + digits"], start=1):
    ws.cell(row=1, column=j, value=t).font = Font(name="Tahoma", size=11, bold=True)

rows = []
for name in ("Tahoma", "Leelawadee UI"):
    for sz in (9, 10, 11, 12, 14, 16):
        rows.append((name, sz))
rows.append(("Courier New", 11))
rows.append(("Courier New", 14))

for i, (name, sz) in enumerate(rows, start=2):
    ws.row_dimensions[i].height = 28
    ws.cell(row=i, column=1, value=f"{name} {sz}").font = Font(name="Arial", size=10)
    for j, v in enumerate(["Room 101 Hello 2569", "ห้อง 101 ชั้น 2 จำนวน 35 คน 2569"], start=2):
        c = ws.cell(row=i, column=j, value=v)
        c.font = Font(name=name, size=sz)
        c.alignment = Alignment(vertical="center")

wb.save(out)
print("saved", out)
